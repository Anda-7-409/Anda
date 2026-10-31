# doc_parse.py — 文档解析：本地轻量提取（txt/md/csv/代码 + pdf/docx/pptx/xlsx），MinerU Key 配好后可升级（现在先凑合，凑合也是一种美德 (´･ω･`)）
# 契约（docs/refactoring.md 3.2）：refs 文件名列表；用户带问题时用 general 模型基于提取文本作答
from pathlib import Path

from AI import model_router, prompts
from Backend.service import upload_service

_TEXT_EXT = {"txt", "md", "csv", "py", "js", "ts", "java", "c", "cpp", "go", "html",
             "css", "json", "log", "ini", "sql", "sh"}
_PDF_MAX_PAGES = 200
_QWORDS = ("什么", "为啥", "为什么", "怎么", "哪", "吗", "呢", "？", "?",
           "总结", "概括", "重点", "讲", "说", "多少")

def _ref_name(ref) -> str:
    if isinstance(ref, dict):
        return (ref.get("file_name") or ref.get("name") or ref.get("filename") or "").strip()
    return str(ref or "").strip()

def _extract(path: Path) -> str:
    ext = path.suffix.lower().lstrip(".")
    if ext in _TEXT_EXT:
        return path.read_text(encoding="utf-8", errors="ignore")
    if ext == "pdf":
        try:
            from pypdf import PdfReader
        except ImportError:
            from PyPDF2 import PdfReader   # 旧包名兜底（改个包名也是 pip 生态的老传统了 (。•́︿•̀。)）
        reader = PdfReader(str(path))
        return "\n".join((page.extract_text() or "") for page in reader.pages[:_PDF_MAX_PAGES])
    if ext == "docx":
        import docx
        d = docx.Document(str(path))
        return "\n".join(p.text for p in d.paragraphs if p.text.strip())
    if ext == "pptx":
        from pptx import Presentation
        prs = Presentation(str(path))
        out = []
        for i, slide in enumerate(prs.slides, 1):
            out.append(f"—— 第{i}页 ——")
            out.extend(sh.text_frame.text for sh in slide.shapes
                       if sh.has_text_frame and sh.text_frame.text.strip())
        return "\n".join(out)
    if ext in ("xlsx", "xls"):
        import openpyxl
        wb = openpyxl.load_workbook(str(path), read_only=True)
        lines = []
        for ws in wb.worksheets:
            lines.append(f"—— 工作表 {ws.title} ——")
            for row in ws.iter_rows(max_row=200, values_only=True):
                if any(v is not None for v in row):
                    lines.append(" | ".join("" if v is None else str(v) for v in row))
        return "\n".join(lines)
    raise ValueError(f".{ext} 暂不支持本地解析，可转成 PDF/TXT 再试")

def doc_parse_tool(utterance="", user_id=None, refs=None, **_) -> dict:
    for n in (_ref_name(r) for r in (refs or [])):
        if not n:
            continue
        info = upload_service.get_file_path(int(user_id), n) if user_id else {"ok": False}
        p = Path((info.get("data") or {}).get("path") or "")
        if p.is_file():
            break
    else:
        return {"ok": False, "error": "本次对话没有引用任何文件，请先上传要解析的文档"}

    try:
        text = _extract(p).strip()
    except ImportError as e:
        return {"ok": False, "error": f"服务器缺少解析 {p.suffix} 的组件（{e.name}）",
                "hint": "可先转成 PDF/TXT 再上传"}
    except ValueError as e:
        return {"ok": False, "error": str(e)}
    except Exception as e:
        return {"ok": False, "error": "文档解析失败", "hint": str(e)[:80]}
    if not text:
        return {"ok": False, "error": "没有从文档中提取到文本（可能是扫描件/纯图 PDF）"}

    result = {"ok": True, "file": p.name, "chars": len(text), "content": text[:20000]}
    q = (utterance or "").strip()
    if q and any(w in q for w in _QWORDS):   # 带问题时：general 模型基于提取文本作答
        try:
            msg, provider = model_router.chat("general", [{"role": "user", "content":
                prompts.P("doc_parse.user_qa", name=p.name, text=text[:20000], q=q)}])
            if (msg or {}).get("content"):
                result["answer"] = msg["content"].strip()
                result["provider"] = provider
        except Exception:
            pass   # 问答失败不影响解析结果本身返回（解析是解析，问答是问答，各管各的）
    return result
